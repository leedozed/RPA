from openpyxl import load_workbook
wb = load_workbook("sample.xlsx")
ws = wb.active

from openpyxl.chart import BarChart, LineChart, Reference
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart()
# bar_chart.add_data(bar_value) #차트 데이터 추가
# ws.add_chart(bar_chart, "E1")

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) #범례에 영어, 수학 넣기
line_chart.title = "성적표"
line_chart.style = 20 #미리 선정된 스타일
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"

ws.add_chart(line_chart, "E1")



wb.save("sampel_chart.xlsx")

